"""Extract catalog products from scanned/uploaded files via LLM (vision / text)."""

from __future__ import annotations

import base64
import csv
import io
import logging
import re
from typing import Any

from fastapi import UploadFile

from services.llm import chat_json

log = logging.getLogger(__name__)

# Shared JSON shape — field meaning depends on business_type (see prompts below).
EXTRACT_JSON_SHAPE = """
Return JSON only:
{
  "items": [
    {
      "name": "string",
      "price": 0,
      "vendor_category": "string or null",
      "portion_size": "string or null",
      "delivery_time_minutes": null,
      "modifiers": []
    }
  ]
}
""".strip()

RESTAURANT_EXTRACT_JSON_SHAPE = """
Return JSON only:
{
  "items": [
    {
      "name": "string",
      "price": 0,
      "vendor_category": "string or null",
      "portion_size": "string or null",
      "delivery_time_minutes": null,
      "modifiers": [
        {
          "group": "protein",
          "options": [{"label": "string", "price_delta": 0}]
        }
      ]
    }
  ]
}
""".strip()

_MODIFIER_GROUP_CANONICAL = {
    "protein": "protein",
    "meat": "protein",
    "choice of protein": "protein",
    "extras": "extras",
    "extra": "extras",
    "sauce": "extras",
    "sauces": "extras",
    "topping": "extras",
    "toppings": "extras",
    "add-on": "extras",
    "addon": "extras",
    "add ons": "extras",
    "add-ons": "extras",
    "size": "size",
    "sizes": "size",
}


def canonical_modifier_group(raw: str) -> str | None:
    key = " ".join((raw or "").strip().lower().split())
    if key in _MODIFIER_GROUP_CANONICAL:
        return _MODIFIER_GROUP_CANONICAL[key]
    for needle, canon in (
        ("protein", "protein"),
        ("meat", "protein"),
        ("sauce", "extras"),
        ("extra", "extras"),
        ("topping", "extras"),
        ("add-on", "extras"),
        ("addon", "extras"),
        ("size", "size"),
    ):
        if needle in key:
            return canon
    return None


# Back-compat alias for internal callers
_canonical_modifier_group = canonical_modifier_group


def _normalize_modifiers(raw_mods: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_mods, list):
        return []
    by_group: dict[str, list[dict[str, Any]]] = {}
    for mod in raw_mods:
        if not isinstance(mod, dict):
            continue
        group = _canonical_modifier_group(str(mod.get("group") or mod.get("name") or ""))
        if not group:
            continue
        options_raw = mod.get("options") or []
        if not isinstance(options_raw, list):
            continue
        bucket = by_group.setdefault(group, [])
        seen = {o["label"].lower() for o in bucket}
        for opt in options_raw:
            if not isinstance(opt, dict):
                continue
            label = str(opt.get("label") or opt.get("name") or "").strip()
            if not label or label.lower() in seen:
                continue
            try:
                delta = float(opt.get("price_delta") if opt.get("price_delta") is not None else opt.get("price") or 0)
            except (TypeError, ValueError):
                delta = 0.0
            bucket.append({"label": label, "price_delta": delta})
            seen.add(label.lower())

    order = ("protein", "extras", "size")
    result = [
        {"group": g, "options": by_group[g]}
        for g in order
        if g in by_group and by_group[g]
    ]
    # Size only when varieties exist
    return [
        row
        for row in result
        if row["group"] != "size" or len(row["options"]) >= 2
    ]


def _extract_system_prompt(business_type: str) -> str:
    bt = (business_type or "Restaurant").strip()
    key = bt.lower()

    if key == "pharmacy":
        return f"""You extract sellable pharmacy products from a price list or catalog.
Business type: Pharmacy.
{EXTRACT_JSON_SHAPE}
Rules:
- price must be a number in Naira (no currency symbol)
- name should include strength when visible (e.g. "Paracetamol 500mg")
- vendor_category is the section label (e.g. Pain relief, Vitamins, OTC)
- portion_size means pack form when visible: Strip, Sachet, Bottle, Box, or pack count — otherwise null
- delivery_time_minutes must always be null for pharmacies
- modifiers must always be []
- skip headers, addresses, license text, and non-product lines
- return at most 40 items
"""

    if key == "shop":
        return f"""You extract sellable shop / retail inventory items from a price list or catalog.
Business type: Shop.
{EXTRACT_JSON_SHAPE}
Rules:
- price must be a number in Naira (no currency symbol)
- vendor_category is the aisle/section label (e.g. Groceries, Household, Electronics)
- portion_size means pack size when visible: Single, Pack, Dozen, Carton, or weight/count — otherwise null
- delivery_time_minutes must always be null for shops
- modifiers must always be []
- skip headers, store info, and non-product lines
- return at most 40 items
"""

    if key == "market":
        return f"""You extract sellable market stall items from a price list or board.
Business type: Market.
{EXTRACT_JSON_SHAPE}
Rules:
- price must be a number in Naira (no currency symbol)
- vendor_category is the produce group (e.g. Produce, Grains, Meat & fish)
- portion_size means sale unit when visible: Small, Medium, Large, Heap, Bag, or weight — otherwise null
- delivery_time_minutes must always be null for markets
- modifiers must always be []
- skip headers and non-product lines
- return at most 40 items
"""

    # Restaurant (default)
    return f"""You extract sellable menu items from a restaurant menu document.
Business type: Restaurant.
{RESTAURANT_EXTRACT_JSON_SHAPE}
Rules:
- price must be a number in Naira (no currency symbol)
- vendor_category is the menu section as written (e.g. Swallow, Grill, Drinks, Desserts, Pastries)
- There is NO Bakery browse tab. Keep pastry/bakery/confectionery section names as vendor_category; they map later into Food, Desserts, or Sides
- portion_size is portion when visible: Small, Medium, Large, Regular — otherwise null
- delivery_time_minutes is prep/ready time in minutes when stated; otherwise null
- Extract per-item customizations into modifiers when listed (choice of protein, extras, sauces, toppings, size)
- modifier.group must be only: "protein", "extras", or "size"
- Put sauces, toppings, and add-ons under "extras" (do not use a separate sauce group)
- Put size varieties (Small/Medium/Large with different prices) under modifiers group "size" only when 2+ sizes are listed for that item
- If only one portion/size is shown, set portion_size and leave size modifiers empty
- Protein as a full dish (e.g. full chicken, grilled fish plate) is its own item — not a modifier
- Protein as an add-on choice for a meal (e.g. goat, beef, fish with swallow) goes in modifiers group "protein"
- price_delta is the add-on cost in Naira (0 if included / no extra charge shown)
- skip headers, phone numbers, addresses, and non-item lines
- return at most 40 items
"""


def _normalize_items(raw: Any, *, business_type: str) -> list[dict[str, Any]]:
    if isinstance(raw, dict):
        items = raw.get("items") or raw.get("products") or []
    elif isinstance(raw, list):
        items = raw
    else:
        items = []

    bt = (business_type or "Restaurant").strip().lower()
    allow_duration = bt == "restaurant"

    out: list[dict[str, Any]] = []
    for row in items:
        if not isinstance(row, dict):
            continue
        name = str(row.get("name") or "").strip()
        if not name:
            continue
        price_raw = row.get("price")
        try:
            if isinstance(price_raw, str):
                price_raw = re.sub(r"[^\d.]", "", price_raw.replace(",", ""))
            price = float(price_raw)
        except (TypeError, ValueError):
            continue
        if price <= 0:
            continue

        delivery_time = None
        if allow_duration:
            delivery = row.get("delivery_time_minutes")
            try:
                delivery_time = int(delivery) if delivery is not None else None
            except (TypeError, ValueError):
                delivery_time = None

        vendor_category = row.get("vendor_category") or row.get("category")
        portion = row.get("portion_size") or row.get("pack_size")
        modifiers = _normalize_modifiers(row.get("modifiers")) if allow_duration else []
        out.append(
            {
                "name": name,
                "price": price,
                "vendor_category": str(vendor_category).strip() if vendor_category else None,
                "portion_size": str(portion).strip() if portion else None,
                "delivery_time": delivery_time,
                "modifiers": modifiers,
            }
        )
    return out


async def _extract_from_text(text: str, business_type: str) -> list[dict[str, Any]]:
    clipped = text.strip()
    if not clipped:
        return []
    if len(clipped) > 24000:
        clipped = clipped[:24000]
    raw = await chat_json(
        system=_extract_system_prompt(business_type),
        user=(
            f"Business type: {business_type}\n"
            "Extract products from this document text:\n\n"
            f"{clipped}"
        ),
        use_vision=False,
    )
    return _normalize_items(raw, business_type=business_type)


async def _extract_from_image(
    data: bytes,
    content_type: str,
    business_type: str,
) -> list[dict[str, Any]]:
    mime = content_type or "image/jpeg"
    b64 = base64.b64encode(data).decode("ascii")
    data_url = f"data:{mime};base64,{b64}"
    raw = await chat_json(
        system=_extract_system_prompt(business_type),
        user=[
            {
                "type": "text",
                "text": (
                    f"Business type: {business_type}\n"
                    "Extract every sellable item with prices from this catalog/menu image."
                ),
            },
            {"type": "image_url", "image_url": {"url": data_url}},
        ],
        use_vision=True,
    )
    return _normalize_items(raw, business_type=business_type)


def _read_csv_text(data: bytes) -> str:
    text = data.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    lines = [", ".join(cell.strip() for cell in row if cell.strip()) for row in reader]
    return "\n".join(line for line in lines if line)


def _read_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:20]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _read_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word documents") from exc
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


async def extract_catalog_items_from_upload(
    file: UploadFile,
    *,
    business_type: str,
) -> list[dict[str, Any]]:
    data = await file.read()
    if not data:
        log.warning("Catalog extract received empty file %s", file.filename)
        return []

    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    log.info(
        "Catalog extract parsing file=%s content_type=%s bytes=%s business_type=%s",
        filename or file.filename,
        content_type or "(none)",
        len(data),
        business_type,
    )

    is_image = content_type.startswith("image/") or filename.endswith(
        (".jpg", ".jpeg", ".png", ".webp", ".gif")
    )
    if is_image:
        log.info("Catalog extract path=image for %s", filename)
        return await _extract_from_image(data, content_type or "image/jpeg", business_type)

    if content_type == "application/pdf" or filename.endswith(".pdf"):
        text = _read_pdf_text(data)
        if len(text.strip()) < 20:
            log.warning("PDF text extraction yielded little content for %s", filename)
        else:
            log.info("Catalog extract path=pdf chars=%s for %s", len(text), filename)
        return await _extract_from_text(text, business_type)

    if filename.endswith((".csv", ".txt")) or content_type in {
        "text/csv",
        "text/plain",
        "application/csv",
    }:
        text = _read_csv_text(data) if filename.endswith(".csv") else data.decode("utf-8", errors="ignore")
        log.info("Catalog extract path=text/csv chars=%s for %s", len(text), filename)
        return await _extract_from_text(text, business_type)

    if filename.endswith((".doc", ".docx")) or "word" in content_type:
        text = _read_docx_text(data)
        log.info("Catalog extract path=docx chars=%s for %s", len(text), filename)
        return await _extract_from_text(text, business_type)

    if filename.endswith((".xls", ".xlsx")) or "excel" in content_type or "spreadsheet" in content_type:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for Excel files") from exc
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets[:3]:
            for row in sheet.iter_rows(max_row=200, values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(", ".join(cells))
        log.info("Catalog extract path=excel lines=%s for %s", len(lines), filename)
        return await _extract_from_text("\n".join(lines), business_type)

    log.info("Catalog extract path=fallback-bytes for %s", filename)
    return await _extract_from_text(data.decode("utf-8", errors="ignore"), business_type)
